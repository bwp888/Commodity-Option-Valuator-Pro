"""
Commodity Option Valuator Pro
=============================

Recommendation Report Excel Export.

Commit 0023
-----------

Exports RecommendationReportPresentation into an Excel workbook.

The exporter is intentionally kept at the output layer.
It does not change recommendation rules, report data,
ranking order, or presentation values.

Workbook
--------

Summary
    Report-level information and recommendation statistics.

Recommendations
    One row per recommendation item.

Author : Simon
Version : 0.6.4
Python : 3.12
"""

from __future__ import annotations

from pathlib import Path
from typing import TYPE_CHECKING

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

if TYPE_CHECKING:
    from core.recommendation_report_presentation import (
        RecommendationReportPresentation,
        RecommendationReportPresentationRow,
    )


# ==========================================================
# Worksheet Names
# ==========================================================

SUMMARY_SHEET_NAME = "Summary"
RECOMMENDATIONS_SHEET_NAME = "Recommendations"


# ==========================================================
# Headers
# ==========================================================

RECOMMENDATION_HEADERS = (
    "Symbol",
    "Action",
    "Action Title",
    "Level",
    "Level Title",
    "Score",
    "Risk Score",
    "Reason",
    "Risk Warning",
    "Action Description",
    "Level Description",
)


# ==========================================================
# Validation
# ==========================================================


def _validate_presentation(
    presentation: RecommendationReportPresentation,
) -> None:
    from core.recommendation_report_presentation import (
        RecommendationReportPresentation,
    )

    if not isinstance(
        presentation,
        RecommendationReportPresentation,
    ):
        raise TypeError(
            "presentation must be a "
            "RecommendationReportPresentation"
        )


# ==========================================================
# Summary Sheet
# ==========================================================


def _write_summary_sheet(
    worksheet: Worksheet,
    presentation: RecommendationReportPresentation,
) -> None:
    worksheet["A1"] = "Report Information"

    summary_rows = (
        ("Title", presentation.title),
        (
            "Generated At",
            presentation.generated_at_text,
        ),
        (
            "Ranked Count",
            presentation.ranked_count,
        ),
        (
            "Total Count",
            presentation.total_count,
        ),
        (
            "Active Count",
            presentation.active_count,
        ),
        (
            "Has Active Recommendation",
            presentation.has_active_recommendation,
        ),
        (
            "Has High Risk",
            presentation.has_high_risk,
        ),
        (
            "Highest Score",
            presentation.highest_score,
        ),
        (
            "Lowest Risk Score",
            presentation.lowest_risk_score,
        ),
        (
            "Top Symbol",
            presentation.top_symbol,
        ),
        (
            "Top Action",
            presentation.top_action,
        ),
        (
            "Top Level",
            presentation.top_level,
        ),
    )

    for row_number, (label, value) in enumerate(
        summary_rows,
        start=2,
    ):
        worksheet.cell(
            row=row_number,
            column=1,
            value=label,
        )
        worksheet.cell(
            row=row_number,
            column=2,
            value=value,
        )

    action_start_row = len(summary_rows) + 4

    worksheet.cell(
        row=action_start_row,
        column=1,
        value="Action Counts",
    )

    for offset, (action, count) in enumerate(
        presentation.action_counts.items(),
        start=1,
    ):
        worksheet.cell(
            row=action_start_row + offset,
            column=1,
            value=action,
        )
        worksheet.cell(
            row=action_start_row + offset,
            column=2,
            value=count,
        )

    level_start_row = (
        action_start_row
        + len(presentation.action_counts)
        + 2
    )

    worksheet.cell(
        row=level_start_row,
        column=1,
        value="Level Counts",
    )

    for offset, (level, count) in enumerate(
        presentation.level_counts.items(),
        start=1,
    ):
        worksheet.cell(
            row=level_start_row + offset,
            column=1,
            value=level,
        )
        worksheet.cell(
            row=level_start_row + offset,
            column=2,
            value=count,
        )

    worksheet.column_dimensions["A"].width = 32
    worksheet.column_dimensions["B"].width = 28


# ==========================================================
# Recommendations Sheet
# ==========================================================


def _row_values(
    row: RecommendationReportPresentationRow,
) -> tuple[object, ...]:
    return (
        row.symbol,
        row.action,
        row.action_title,
        row.level,
        row.level_title,
        row.score,
        row.risk_score,
        row.reason,
        row.risk_warning,
        row.action_description,
        row.level_description,
    )


def _write_recommendations_sheet(
    worksheet: Worksheet,
    presentation: RecommendationReportPresentation,
) -> None:
    worksheet.append(RECOMMENDATION_HEADERS)

    for row in presentation.rows:
        worksheet.append(
            _row_values(row)
        )

    widths = {
        "A": 18,
        "B": 12,
        "C": 16,
        "D": 10,
        "E": 14,
        "F": 12,
        "G": 14,
        "H": 40,
        "I": 40,
        "J": 40,
        "K": 40,
    }

    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width


# ==========================================================
# Public Export Function
# ==========================================================


def export_recommendation_report(
    presentation: RecommendationReportPresentation,
    output_path: str | Path,
) -> Path:
    """
    Export a RecommendationReportPresentation to Excel.

    Parameters
    ----------
    presentation:
        The already-presented recommendation report.

    output_path:
        Target .xlsx file path.

    Returns
    -------
    Path
        The resolved output path.

    Raises
    ------
    TypeError
        If presentation is not a
        RecommendationReportPresentation.

    ValueError
        If output_path does not use the .xlsx extension.
    """

    _validate_presentation(presentation)

    path = Path(output_path)

    if path.suffix.lower() != ".xlsx":
        raise ValueError(
            "output_path must use the .xlsx extension"
        )

    path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    workbook = Workbook()

    summary_sheet = workbook.active
    summary_sheet.title = SUMMARY_SHEET_NAME

    recommendations_sheet = workbook.create_sheet(
        RECOMMENDATIONS_SHEET_NAME
    )

    _write_summary_sheet(
        summary_sheet,
        presentation,
    )

    _write_recommendations_sheet(
        recommendations_sheet,
        presentation,
    )

    workbook.save(path)

    return path.resolve()


__all__ = [
    "SUMMARY_SHEET_NAME",
    "RECOMMENDATIONS_SHEET_NAME",
    "RECOMMENDATION_HEADERS",
    "export_recommendation_report",
]