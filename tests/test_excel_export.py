"""
Commodity Option Valuator Pro
=============================

Tests for Recommendation Report Excel Export.

Commit 0023
-----------

Author : Simon
Version : 0.6.4
"""

from __future__ import annotations

from datetime import datetime

import pytest
from openpyxl import load_workbook

from core.recommendation_presentation import (
    PresentationAction,
    PresentationLevel,
    RecommendationPresentation,
    RecommendationPresentationResult,
)
from core.recommendation_report import (
    RecommendationReport,
)
from core.recommendation_summary import (
    RecommendationSummary,
)
from reports.excel_export import (
    RECOMMENDATION_HEADERS,
    RECOMMENDATIONS_SHEET_NAME,
    SUMMARY_SHEET_NAME,
    export_recommendation_report,
)


# ==========================================================
# Test Data
# ==========================================================


GENERATED_AT = datetime(
    2026,
    8,
    19,
    10,
    30,
    0,
)


def make_summary(
    total_count: int = 2,
    buy_count: int = 1,
    sell_count: int = 0,
    watch_count: int = 1,
    reject_count: int = 0,
) -> RecommendationSummary:
    """Create deterministic summary data."""

    return RecommendationSummary(
        total_count=total_count,
        buy_count=buy_count,
        sell_count=sell_count,
        watch_count=watch_count,
        reject_count=reject_count,
        level_a_count=buy_count,
        level_b_count=watch_count,
        level_c_count=0,
        level_d_count=reject_count,
        highest_score=8.25 if total_count else None,
        lowest_risk_score=3.5 if total_count else None,
        top=None,
    )


def make_item(
    symbol: str = "TEST-C-100",
    action: PresentationAction = PresentationAction.BUY,
    level: PresentationLevel = PresentationLevel.A,
    score: float = 8.25,
    risk_score: float = 3.5,
    risk_warning: str | None = "test warning",
) -> RecommendationPresentation:
    """Create deterministic presentation item."""

    return RecommendationPresentation(
        symbol=symbol,
        action=action,
        action_title=action.value,
        action_description="test action",
        level=level,
        level_title=level.value,
        level_description="test level",
        score=score,
        risk_score=risk_score,
        reason="test reason",
        risk_warning=risk_warning,
    )


def make_presentation_result(
    items: tuple[
        RecommendationPresentation,
        ...,
    ],
) -> RecommendationPresentationResult:
    """Create deterministic presentation result."""

    summary = make_summary(
        total_count=len(items),
        buy_count=sum(
            item.action == PresentationAction.BUY
            for item in items
        ),
        sell_count=sum(
            item.action == PresentationAction.SELL
            for item in items
        ),
        watch_count=sum(
            item.action == PresentationAction.WATCH
            for item in items
        ),
        reject_count=sum(
            item.action == PresentationAction.REJECT
            for item in items
        ),
    )

    return RecommendationPresentationResult(
        items=items,
        summary=summary,
        ranked_count=len(items),
    )


def make_report(
    items: tuple[
        RecommendationPresentation,
        ...,
    ] | None = None,
) -> RecommendationReport:
    """Create a deterministic RecommendationReport."""

    if items is None:
        items = (
            make_item(),
            make_item(
                symbol="TEST-P-110",
                action=PresentationAction.WATCH,
                level=PresentationLevel.B,
                score=7.50,
                risk_score=5.00,
                risk_warning=None,
            ),
        )

    result = make_presentation_result(
        items
    )

    return RecommendationReport.from_presentation(
        result,
        title="Test Recommendation Report",
        generated_at=GENERATED_AT,
    )


def make_presentation():
    """Create the final report presentation."""

    report = make_report()

    from core.recommendation_report_presentation import (
        RecommendationReportPresenter,
    )

    return RecommendationReportPresenter.present(
        report
    )


# ==========================================================
# Export File
# ==========================================================


def test_export_creates_xlsx_file(tmp_path):
    """Export should create the requested XLSX file."""

    presentation = make_presentation()

    output_path = (
        tmp_path
        / "recommendation_report.xlsx"
    )

    result = export_recommendation_report(
        presentation,
        output_path,
    )

    assert result == output_path.resolve()
    assert output_path.exists()
    assert output_path.is_file()


# ==========================================================
# Workbook Structure
# ==========================================================


def test_export_creates_expected_sheets(
    tmp_path,
):
    """Workbook should contain the two defined sheets."""

    presentation = make_presentation()

    output_path = (
        tmp_path
        / "recommendation_report.xlsx"
    )

    export_recommendation_report(
        presentation,
        output_path,
    )

    workbook = load_workbook(
        output_path,
        read_only=True,
        data_only=True,
    )

    assert workbook.sheetnames == [
        SUMMARY_SHEET_NAME,
        RECOMMENDATIONS_SHEET_NAME,
    ]


# ==========================================================
# Summary
# ==========================================================


def test_summary_sheet_contains_report_information(
    tmp_path,
):
    """Summary sheet should preserve report-level values."""

    presentation = make_presentation()

    output_path = (
        tmp_path
        / "recommendation_report.xlsx"
    )

    export_recommendation_report(
        presentation,
        output_path,
    )

    workbook = load_workbook(
        output_path,
        data_only=True,
    )

    worksheet = workbook[
        SUMMARY_SHEET_NAME
    ]

    assert worksheet["A1"].value == (
        "Report Information"
    )

    assert worksheet["A2"].value == "Title"
    assert worksheet["B2"].value == (
        "Test Recommendation Report"
    )

    assert worksheet["A3"].value == "Generated At"
    assert worksheet["B3"].value == (
        "2026-08-19 10:30:00"
    )

    assert worksheet["A4"].value == "Ranked Count"
    assert worksheet["B4"].value == 2

    assert worksheet["A5"].value == "Total Count"
    assert worksheet["B5"].value == 2

    assert worksheet["A6"].value == "Active Count"
    assert worksheet["B6"].value == 1


def test_summary_sheet_contains_flags_and_scores(
    tmp_path,
):
    """Summary should preserve flags and score information."""

    presentation = make_presentation()

    output_path = (
        tmp_path
        / "recommendation_report.xlsx"
    )

    export_recommendation_report(
        presentation,
        output_path,
    )

    workbook = load_workbook(
        output_path,
        data_only=True,
    )

    worksheet = workbook[
        SUMMARY_SHEET_NAME
    ]

    values = {
        worksheet.cell(
            row=row,
            column=1,
        ).value:
        worksheet.cell(
            row=row,
            column=2,
        ).value
        for row in range(1, 30)
    }

    assert values[
        "Has Active Recommendation"
    ] is True

    assert values[
        "Has High Risk"
    ] is True

    assert values[
        "Highest Score"
    ] == 8.25

    assert values[
        "Lowest Risk Score"
    ] == 3.5


def test_summary_sheet_contains_top_information(
    tmp_path,
):
    """Summary should preserve top recommendation information."""

    presentation = make_presentation()

    output_path = (
        tmp_path
        / "recommendation_report.xlsx"
    )

    export_recommendation_report(
        presentation,
        output_path,
    )

    workbook = load_workbook(
        output_path,
        data_only=True,
    )

    worksheet = workbook[
        SUMMARY_SHEET_NAME
    ]

    values = {
        worksheet.cell(
            row=row,
            column=1,
        ).value:
        worksheet.cell(
            row=row,
            column=2,
        ).value
        for row in range(1, 30)
    }

    assert values["Top Symbol"] == "TEST-C-100"
    assert values["Top Action"] == "BUY"
    assert values["Top Level"] == "A"


# ==========================================================
# Recommendation Sheet
# ==========================================================


def test_recommendations_sheet_contains_expected_headers(
    tmp_path,
):
    """Recommendation sheet should contain the defined headers."""

    presentation = make_presentation()

    output_path = (
        tmp_path
        / "recommendation_report.xlsx"
    )

    export_recommendation_report(
        presentation,
        output_path,
    )

    workbook = load_workbook(
        output_path,
        data_only=True,
    )

    worksheet = workbook[
        RECOMMENDATIONS_SHEET_NAME
    ]

    headers = tuple(
        cell.value
        for cell in worksheet[1]
    )

    assert headers == RECOMMENDATION_HEADERS


def test_recommendations_sheet_preserves_row_order(
    tmp_path,
):
    """Recommendation order should be unchanged."""

    presentation = make_presentation()

    output_path = (
        tmp_path
        / "recommendation_report.xlsx"
    )

    export_recommendation_report(
        presentation,
        output_path,
    )

    workbook = load_workbook(
        output_path,
        data_only=True,
    )

    worksheet = workbook[
        RECOMMENDATIONS_SHEET_NAME
    ]

    assert worksheet["A2"].value == "TEST-C-100"
    assert worksheet["A3"].value == "TEST-P-110"


def test_recommendations_sheet_preserves_values(
    tmp_path,
):
    """Recommendation fields should be exported unchanged."""

    presentation = make_presentation()

    output_path = (
        tmp_path
        / "recommendation_report.xlsx"
    )

    export_recommendation_report(
        presentation,
        output_path,
    )

    workbook = load_workbook(
        output_path,
        data_only=True,
    )

    worksheet = workbook[
        RECOMMENDATIONS_SHEET_NAME
    ]

    assert worksheet["A2"].value == "TEST-C-100"
    assert worksheet["B2"].value == "BUY"
    assert worksheet["C2"].value == "BUY"
    assert worksheet["D2"].value == "A"
    assert worksheet["E2"].value == "A"
    assert worksheet["F2"].value == 8.25
    assert worksheet["G2"].value == 3.5
    assert worksheet["H2"].value == "test reason"
    assert worksheet["I2"].value == "test warning"
    assert worksheet["J2"].value == "test action"
    assert worksheet["K2"].value == "test level"


def test_recommendations_sheet_preserves_empty_warning(
    tmp_path,
):
    """A missing risk warning should remain empty in Excel."""

    presentation = make_presentation()

    output_path = (
        tmp_path
        / "recommendation_report.xlsx"
    )

    export_recommendation_report(
        presentation,
        output_path,
    )

    workbook = load_workbook(
        output_path,
        data_only=True,
    )

    worksheet = workbook[
        RECOMMENDATIONS_SHEET_NAME
    ]

    assert worksheet["I3"].value is None


# ==========================================================
# Output Path
# ==========================================================


def test_export_creates_parent_directory(
    tmp_path,
):
    """Exporter should create missing parent directories."""

    presentation = make_presentation()

    output_path = (
        tmp_path
        / "nested"
        / "directory"
        / "report.xlsx"
    )

    export_recommendation_report(
        presentation,
        output_path,
    )

    assert output_path.exists()


def test_export_rejects_invalid_presentation(
    tmp_path,
):
    """Invalid presentation input should be rejected."""

    output_path = (
        tmp_path
        / "report.xlsx"
    )

    with pytest.raises(
        TypeError,
        match=(
            "presentation must be a "
            "RecommendationReportPresentation"
        ),
    ):
        export_recommendation_report(
            object(),
            output_path,
        )


def test_export_rejects_non_xlsx_path(
    tmp_path,
):
    """Only XLSX output should be accepted."""

    presentation = make_presentation()

    output_path = (
        tmp_path
        / "report.csv"
    )

    with pytest.raises(
        ValueError,
        match=(
            "output_path must use the .xlsx extension"
        ),
    ):
        export_recommendation_report(
            presentation,
            output_path,
        )